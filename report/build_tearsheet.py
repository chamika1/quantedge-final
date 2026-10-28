"""One-page executive tear sheet, built from results/results.json."""
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
R = json.loads((ROOT / "results" / "results.json").read_text())
OUT = ROOT / "report"
NAVY = RGBColor(0x1F, 0x4E, 0x79)

doc = Document()
s = doc.sections[0]
s.top_margin, s.bottom_margin = Inches(0.45), Inches(0.4)
s.left_margin, s.right_margin = Inches(0.6), Inches(0.6)
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(9.5)
st.paragraph_format.space_after = Pt(3)
st.paragraph_format.line_spacing = 1.05


def para(text, bold=False, size=9.5, align=None, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)
    if color:
        r.font.color.rgb = color
    if align:
        p.alignment = align
    return p


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, txt in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = str(txt)
        for par in c.paragraphs:
            for run in par.runs:
                run.bold = True
                run.font.size = Pt(8.8)
    for row in rows:
        cells = t.add_row().cells
        for i, txt in enumerate(row):
            cells[i].text = str(txt)
            for par in cells[i].paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8.8)
    doc.add_paragraph()
    return t


def pct(x, dp=1):
    return f"{x * 100:.{dp}f}%"


def usd(x):
    return f"USD {x:,.0f}"


W = R["meta"]["weights"]
K, KE = R["kpis_full_history"], R["kpis_estimation_window"]
RT = {r["level"]: r for r in R["risk_table_estimation"]}
REP = {r["scenario"]: r for r in R["stress_replays"]}
HYP = R["stress_hypothetical"]
BT = R["backtests"]

p = para("TEAR SHEET - SERENDIB UNIVERSITY ENDOWMENT FUND", bold=True, size=15,
         align=WD_ALIGN_PARAGRAPH.CENTER)
p.runs[0].font.color.rgb = NAVY
para("Quant Edge 1.0 Final Challenge | Team n0T4ounD, Sabaragamuwa University of "
     "Sri Lanka | USD 10,000,000 - conservative mandate | October 2026",
     size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
para("RECOMMENDATION: A minimum-variance portfolio with extra rules (constraints). "
     "Protect the capital first, keep every asset class in the portfolio, and use "
     "the risk numbers from the real data, not from the normal distribution.",
     bold=True, size=10)

table(["Weights", "SPY", "EFA", "EEM", "AGG", "GLD", "XLE", "XLK"],
      [["", pct(W["SPY"]), pct(W["EFA"]), pct(W["EEM"]), pct(W["AGG"]),
        pct(W["GLD"]), pct(W["XLE"]), pct(W["XLK"])],
       ["USD m", "1.33", "0.55", "0.50", "4.00", "2.00", "1.12", "0.50"]])

table(["Performance (today's weights)", "CAGR", "Volatility", "Sharpe", "Sortino",
       "Beta", "Max drawdown"],
      [["Full history replay 2007-2026", pct(K["CAGR"]), pct(K["AnnVol"]),
        f"{K['Sharpe']:.2f}", f"{K['Sortino']:.2f}", f"{K['Beta_vs_SPY']:.2f}",
        pct(K["MaxDrawdown"])],
       ["Recent window 2021-2026", pct(KE["CAGR"]), pct(KE["AnnVol"]),
        f"{KE['Sharpe']:.2f}", f"{KE['Sortino']:.2f}", f"{KE['Beta_vs_SPY']:.2f}",
        pct(KE["MaxDrawdown"])]])

table(["1-day risk (recent window)", "VaR Historical", "VaR Parametric",
       "VaR EVT", "ES Historical", "ES in USD"],
      [[f"95%", pct(RT[0.95]["VaR_historical"]), pct(RT[0.95]["VaR_parametric"]),
        pct(RT[0.95]["VaR_evt"]), pct(RT[0.95]["ES_historical"]),
        usd(RT[0.95]["ES_historical"] * 10_000_000)],
       [f"99%", pct(RT[0.99]["VaR_historical"]), pct(RT[0.99]["VaR_parametric"]),
        pct(RT[0.99]["VaR_evt"]), pct(RT[0.99]["ES_historical"]),
        usd(RT[0.99]["ES_historical"] * 10_000_000)]])

para(f"WHY THE GAP MATTERS: at 99% the Gaussian model says {pct(RT[0.99]['VaR_parametric'])} "
     f"but the fund's own data say {pct(RT[0.99]['VaR_historical'])}. Returns are "
     f"fat-tailed (excess kurtosis {K['Kurtosis']:.0f}); EVT tail index xi = "
     f"{R['evt']['xi_shape']:.2f} > 0 confirms it. We size risk on the data, not the curve.",
     size=9.5)

table(["Stress replay (today's weights)", "2008 GFC", "2020 COVID crash",
       "2022 rate-hike year", "Designed stagflation scenario"],
      [["Portfolio impact", pct(REP["2008 Global Financial Crisis"]["portfolio_return"]),
        pct(REP["2020 COVID-19 Crash"]["portfolio_return"]),
        pct(REP["2022 Rate-Hike Bear Market"]["portfolio_return"]),
        pct(HYP["deterministic_portfolio_impact"])],
       ["Loss on USD 10m", usd(REP["2008 Global Financial Crisis"]["portfolio_loss_usd"]),
        usd(REP["2020 COVID-19 Crash"]["portfolio_loss_usd"]),
        usd(REP["2022 Rate-Hike Bear Market"]["portfolio_loss_usd"]),
        usd(HYP["deterministic_loss_usd"])]])

para(f"BACKTEST RESULT (19 years, rolling 250-day VaR): the 95% models pass the Kupiec "
     f"test. The 99% models do not pass (historical: {BT['historical_99']['exceptions']} "
     f"exceptions against {BT['historical_99']['kupiec']['expected']:.0f} expected) and "
     "the exceptions come in groups in crises (Christoffersen independence p < 0.001). "
     "Our limits: 99% ES below 3% of the fund (1-day), stress loss below 15%, "
     "rebalance every quarter, re-run the backtests every month.", size=9.5)
para("Full methodology, formulas, limitations and AI-use disclosure: technical report "
     "and Streamlit app (Assumptions & Limitations page). Data: Yahoo Finance adjusted "
     "closes 2007-2026; FRED DTB3 risk-free. Full-history figures are hypothetical "
     "replays of today's weights, before fees. Past performance does not predict "
     "future results.", size=8.5)

doc.save(str(OUT / "QuantEdge_Tear_Sheet.docx"))
print("saved tear sheet")
